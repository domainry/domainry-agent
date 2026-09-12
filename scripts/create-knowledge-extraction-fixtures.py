#!/usr/bin/env python3
"""Generate small synthetic file-format fixtures for the knowledge parser tests."""
from pathlib import Path
import hashlib
import json
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1] / "integration/testdata/knowledge-extraction"
root.mkdir(parents=True, exist_ok=True)
paragraphs = [
    "K06 SYNTHETIC EXTRACTION ACCEPTANCE",
    "This fixture contains no real personal or business information.",
    "Customer: Qinghe Fixture",
    "Invoice Amount: 9007199254740993.25 CNY",
    "Signed Date: 2026-09-11",
    "Approval: true",
    "The contact email is intentionally absent.",
]
rows = [["Item", "Quantity", "Unit Price"], ["Pencil", "2", "1.20"], ["Paper", "3", "12.50"]]
pdf = root / "synthetic-extraction.pdf"
c = canvas.Canvas(str(pdf), pagesize=A4, invariant=1)
c.setTitle("K06 Synthetic Extraction Acceptance")
c.setAuthor("Domainry acceptance fixture")
y = A4[1] - 55
for line in paragraphs + [""] + [" | ".join(row) for row in rows]:
    c.setFont("Helvetica", 11)
    c.drawString(45, y, line)
    y -= 20
c.save()
assert "9007199254740993.25" in "\n".join(page.extract_text() for page in PdfReader(pdf).pages)

doc = Document()
for line in paragraphs:
    doc.add_paragraph(line)
table = doc.add_table(rows=0, cols=3)
table.style = "Table Grid"
for row in rows:
    for cell, value in zip(table.add_row().cells, row):
        cell.text = value
word = root / "synthetic-extraction.docx"
doc.save(word)
check = Document(word)
assert "9007199254740993.25" in "\n".join(p.text for p in check.paragraphs)
assert [[c.text for c in r.cells] for r in check.tables[0].rows] == rows

book = Workbook()
sheet = book.active
sheet.title = "Synthetic Invoice"
for line in paragraphs:
    key, sep, value = line.partition(": ")
    sheet.append([key, value] if sep else [line])
sheet.append([])
for row in rows:
    sheet.append(row)
for column in ["A", "B", "C"]:
    sheet.column_dimensions[column].width = 34
sheet.freeze_panes = "A10"
excel = root / "synthetic-extraction.xlsx"
book.save(excel)
check = load_workbook(excel, data_only=True)
assert check.active["B4"].value == "9007199254740993.25 CNY"
assert check.active["A10"].value == "Pencil"

manifest = {"synthetic": True, "expected": {"customer": "Qinghe Fixture", "amount": "9007199254740993.25", "signed_date": "2026-09-11", "approval": "true", "email": None, "rows": rows[1:]}, "files": []}
for file in [pdf, word, excel]:
    manifest["files"].append({"filename": file.name, "bytes": file.stat().st_size, "sha256": hashlib.sha256(file.read_bytes()).hexdigest()})
(root / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
print("Created and reopened three synthetic PDF/Word/Excel fixtures.")
